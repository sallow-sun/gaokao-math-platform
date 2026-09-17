"""Build the downloadable Word rules from the maintained Markdown source.
Requires python-docx; does not connect to any server or modify old rules.
"""
from pathlib import Path
import re
import shutil
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
docs = root / 'frontend/public/docs'
source = docs / '题目上传规则1.4.md'
text = source.read_text(encoding='utf-8')

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(5)
for style in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[style].font.name = 'Times New Roman'
    doc.styles[style].element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
lines = text.splitlines()
code = False
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith('```'):
        code = not code
    elif code:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [s.strip() for s in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', s) for s in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Table Grid'
        for j, cells in enumerate(rows):
            target = table.rows[0].cells if j == 0 else table.add_row().cells
            for cell, value in zip(target, cells):
                cell.text = value.replace('`', '')
            if j == 0:
                repeat = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
        continue
    elif line.startswith('# '):
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], 2)
    elif line.strip():
        doc.add_paragraph(line.replace('`', '').replace('**', ''))
    i += 1
output = docs / '题目上传规则1.4.docx'
doc.save(output)
if (root / 'TEST').is_dir():
    shutil.copy2(output, root / 'TEST/题目上传规则与模板1.4.docx')
print(output)
